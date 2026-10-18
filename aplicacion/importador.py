"""Lectura por texto, conservación de faltantes y trazabilidad de celda. Sin imputaciones."""
from __future__ import annotations
import csv
import hashlib
import html
import io
import json
import math
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, time, timedelta
from pathlib import Path
import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
VACIOS = {'': 'vacio', '-': 'sin_respuesta', 'n/a': 'no_aplica', 'na': 'no_aplica',
          'no aplica': 'no_aplica', 'desconocido': 'desconocido', 'por definir': 'por_definir'}
MESES = {'ENE-FEB': 1, 'MAR-ABR': 3, 'MAY-JUN': 5, 'JUL-AGO': 7, 'SEP-OCT': 9, 'NOV-DIC': 11}

def texto(v):
    return re.sub(r'\s+', ' ', html.unescape(re.sub('<[^>]*>', ' ', str(v if v is not None else '')))).strip()

def canon(v):
    return ''.join(c for c in unicodedata.normalize('NFKD', texto(v).lower()) if not unicodedata.combining(c))

def pregunta_normalizada(q):
    q = canon(q)
    q = re.sub(r'^(?:avance|bimestre|semestral)\s*-?\s*\d+%\s*-?\s*', '', q)
    return re.sub(r'\[p\d+\]', '[referencia]', q)

def privada(q, d):
    s = canon(q + ' ' + d)
    return bool(re.search(r'nombre|correo|email|e-mail|telefono|celular|whatsapp|curp|\brfc\b|\bnss\b|secretario general|titulo protocolario', s)) or canon(q) in {'link', 'date_modified', 'puesto'}

def categoria(q):
    s = canon(q)
    for cat, patron in [
        ('rotacion', 'bajas|rotacion'), ('reclutamiento', 'contrata|recluta|candidatos|vacantes|plazas'),
        ('ausentismo', 'falta|incapacidad|accidente'), ('capacitacion', 'capacitacion|perfiles.*desarrollando'),
        ('relaciones', 'sindic|negociacion|incremento|conflicto|seguridad'),
        ('salarios', 'salario|salarial'), ('beneficios', 'prestacion|comedor|transporte|sgmm|bono|seguro|uniforme|vacacion'),
        ('contexto', 'empleados|industria|sector|municipio|tier|proveeduria|generacion|mujeres|hombres|capital|operando')]:
        if re.search(patron, s): return cat
    return 'otros'

def sugerir(q, d):
    s, dim = canon(q), canon(d)
    metric, unit, group = '', 'texto', ''
    rules = [
        ('numero de empleados actual de personal operativo', 'plantilla_operativa', 'personas'),
        ('numero de empleados personal administrativo', 'plantilla_administrativa', 'personas'),
        ('otro tipo de personal que labora', 'plantilla_otros', 'personas'),
        ('total de contrataciones', 'contrataciones', 'personas'),
        ('contrataciones de la pregunta anterior', 'contrataciones_desglose', 'personas'),
        ('capture el numero de bajas', 'bajas', 'personas'),
        ('desglose las bajas del bimestre', 'bajas_mes', 'personas'),
        ('desglose de acuerdo al tipo de empleado', 'bajas_tipo', 'personas'),
        ('desglose el numero de bajas de acuerdo', 'bajas_motivo', 'personas'),
        ('que antiguedad tenian los empleados', 'bajas_antiguedad', 'porcentaje'),
        ('que edad tenian los empleados', 'bajas_edad', 'porcentaje'),
        ('total de faltas', 'faltas', 'eventos'),
        ('total de casos de incapacidad', 'incapacidades', 'eventos'),
        ('desglose el numero de incapacidades', 'incapacidades_tipo', 'eventos'),
        ('total de horas de capacitacion', 'capacitacion_horas', 'horas'),
        ('tarda en promedio para reclutar', 'tiempo_contratacion', 'dias'),
        ('numero de plazas operativas y plazas administrativas', 'vacantes_proximas', 'personas'),
        ('estatus actual de su negociacion', 'negociacion_estatus', 'texto'),
        ('comparta el porcentaje de incrementos', 'incremento_otorgado', 'porcentaje'),
        ('que porcentaje de incremento salarial', 'incremento_planeado', 'porcentaje'),
        ('mes o los meses cuando realizan', 'incremento_meses', 'seleccion'),
        ('mes o los meses cuando estiman', 'incremento_meses_previstos', 'seleccion'),
        ('principal medio que utilizan para reclutar', 'fuentes_uso', 'escala_0_10'),
        ('principales fuentes de reclutamiento', 'fuentes_seleccion', 'seleccion'),
        ('que proporcion hay de mujeres y hombres', 'genero', 'porcentaje'),
        ('distribuido su personal de acuerdo a su edad', 'edad_plantilla', 'porcentaje'),
        ('motivos por los cuales son rechazados', 'rechazo_motivo', 'porcentaje'),
        ('principal industria o sector comercial', 'sector', 'texto'),
        ('nivel de proveeduria automotriz', 'tier', 'seleccion'),
        ('costo promedio mensual por persona de su servicio de transporte', 'costo_transporte', 'mxn_persona_mes'),
        ('costo promedio por platillo', 'costo_comedor', 'mxn_platillo'),
        ('sgmm - ¿cual es el costo', 'costo_sgmm', 'mxn_anio'),
    ]
    for match, m, u in rules:
        if match in s:
            metric, unit = m, u
            break
    if s.startswith('municipio') or 'seleccione el municipio' in s: metric, unit = 'municipio', 'texto'
    if 'horario' in dim and ('entrada' in dim or 'salida' in dim): unit = 'hora_del_dia'
    if not metric:
        metric = 'pregunta_' + hashlib.sha256(pregunta_normalizada(q).encode()).hexdigest()[:12]
        if 'porcentaje' in dim or dim.startswith('%'): unit = 'porcentaje'
        elif 'solo admite numeros' in s or 'solo admite valores numericos' in s or 'solo acepta valores numericos' in s: unit = 'numero'
    if metric in {'genero', 'edad_plantilla', 'rechazo_motivo', 'bajas_antiguedad', 'bajas_edad'}: group = metric
    if metric.startswith('plantilla_') or metric in {'contrataciones', 'bajas', 'faltas', 'incapacidades', 'sector', 'municipio', 'negociacion_estatus', 'costo_transporte'}:
        dim = 'total' if dim in {'open-ended response','response',''} else dim
    elif metric in {'fuentes_uso','incremento_otorgado','incremento_planeado','incremento_meses','incremento_meses_previstos'}:
        # Conserva el contexto: operativo/técnico y otorgado/planeado no son intercambiables.
        dim = hashlib.sha256(pregunta_normalizada(q).encode()).hexdigest()[:8] + ':' + dim
    return metric, dim or 'total', unit, group

def columnas(fila1, fila2, overrides=None):
    resultado, pregunta = [], ''
    overrides = overrides or {}
    for i, (a,b) in enumerate(zip(fila1, fila2)):
        if a is None and b is None:
            resultado.append({'columna': i+1, 'accion': 'sin_encabezado'})
            pregunta = ''
            continue
        if a and not str(a).startswith('Unnamed:'): pregunta = texto(a)
        d = texto(b)
        if not pregunta:
            resultado.append({'columna': i+1, 'accion': 'sin_encabezado'})
            continue
        key = hashlib.sha256((pregunta_normalizada(pregunta)+'|'+canon(d)).encode()).hexdigest()[:20]
        m, dim, u, grupo = sugerir(pregunta,d)
        entry = {'clave': key, 'pregunta': pregunta, 'subopcion': d, 'metrica': m, 'dimension': dim,
                 'unidad': u, 'grupo_100': grupo, 'categoria': categoria(pregunta), 'accion': 'incluir'}
        entry.update(overrides.get(key, {}))
        entry['columna'] = i+1
        # La exclusión de PII no se puede desactivar mediante un diccionario.
        if privada(pregunta,d): entry['accion'] = 'descartar_pii'
        if canon(pregunta)=='compania': entry['accion'] = 'identificador'
        resultado.append(entry)
    counts = Counter(c.get('clave') for c in resultado if c['accion']=='incluir')
    for c in resultado:
        if c['accion']=='incluir' and counts[c['clave']]>1: c['duplicada'] = True
    return resultado

def limpiar(v, unidad, formato=''):
    if v is None: return None, None, 'vacio'
    if isinstance(v, bool): return float(v), None, 'valido'
    if isinstance(v, (datetime,time,timedelta)):
        if unidad not in {'horas','hora_del_dia'}: return None, None, 'fecha_no_esperada'
        hours = v.total_seconds()/3600 if isinstance(v,timedelta) else v.hour+v.minute/60+v.second/3600
        return hours, None, 'hora_convertida'
    raw = texto(v)
    n = canon(raw)
    if n in VACIOS: return None, None, VACIOS[n]
    if raw.startswith(('=', '+', '@')) or (raw.startswith('-') and not re.fullmatch(r'-\d+(\.\d+)?',raw)):
        return None, None, 'formula_bloqueada'
    if re.search(r'[\w.+-]+@[\w.-]+\.\w+|(?:\+?\d[ ()-]*){10,}',raw) and not isinstance(v,(float,int)):
        return None, None, 'pii_en_texto_descartada'
    if re.search(r'(mas del?|menos del?|>|<|entre)\s*(?:el\s*)?\$?\d',n): return None, raw[:500], 'intervalo_no_puntual'
    try:
        if isinstance(v,(int,float)): num=float(v)
        else:
            number=raw.rstrip('%').strip()
            if re.fullmatch(r'\d{1,3}(,\d{3})+(\.\d+)?',number): number=number.replace(',','')
            if not re.fullmatch(r'-?\d+(\.\d+)?',number): raise ValueError()
            num=float(number)
        if not math.isfinite(num): return None,None,'no_finito'
        status='valido'
        if unidad=='proporcion':
            if not 0<=num<=1:return None,raw,'proporcion_fuera_rango'
            num*=100;status='proporcion_declarada_convertida'
        if unidad=='porcentaje' and '%' in formato: num*=100; status='porcentaje_excel_convertido'
        if unidad=='hora_del_dia':
            if 0<num<1: num*=24; status='fraccion_dia_convertida'
            if not 0<=num<24: return None,raw,'hora_fuera_rango'
        if unidad=='horas' and ('h' in formato.lower() or '[h]' in formato.lower()): num*=24; status='fraccion_dia_convertida'
        if unidad=='porcentaje' and not 0<=num<=100: return None,raw,'porcentaje_fuera_rango'
        if unidad=='escala_0_10' and not 0<=num<=10: return None,raw,'escala_fuera_rango'
        if unidad in {'personas','eventos'} and (num<0 or not num.is_integer()): return None,raw,'conteo_invalido'
        if unidad not in {'texto','hora_del_dia','seleccion'} and num<0: return None,raw,'negativo'
        return num,None,status
    except ValueError:
        if unidad in {'texto','seleccion'}: return None,raw[:1000],'texto'
        return None,raw[:500],'texto_no_numerico'

def periodo(label):
    m=re.fullmatch(r'(ENE-FEB|MAR-ABR|MAY-JUN|JUL-AGO|SEP-OCT|NOV-DIC)[_ ](20\d{2})',label.upper().strip())
    if not m: raise ValueError('Nombre de hoja inválido: use JUL-AGO_2025, por ejemplo.')
    return f'{m[2]}-{MESES[m[1]]:02d}',m[0].replace(' ','_')

def leer_archivo(data: bytes, nombre: str, overrides=None):
    if len(data)>15*1024*1024: raise ValueError('El límite es 15 MB.')
    overrides=overrides or {}
    sheets=[]
    if nombre.lower().endswith('.xlsx'):
        try:
            z=zipfile.ZipFile(io.BytesIO(data))
            if sum(i.file_size for i in z.infolist())>100*1024*1024: raise ValueError('Excel descomprimido mayor a 100 MB.')
            if any('vbaProject' in i.filename or 'externalLinks/' in i.filename for i in z.infolist()): raise ValueError('No se aceptan macros ni vínculos externos.')
            w=openpyxl.load_workbook(io.BytesIO(data),read_only=True,data_only=False,keep_links=False)
            if len(w.worksheets)>24: raise ValueError('Máximo 24 hojas.')
            if sum((s.max_row or 0)*(s.max_column or 0) for s in w)>1_000_000: raise ValueError('Máximo un millón de celdas por archivo.')
            for s in w:
                if s.max_column>2000 or s.max_row>1000: raise ValueError('Máximo 2,000 columnas y 1,000 filas por hoja.')
                sheets.append((s.title,list(s.iter_rows())))
        except (zipfile.BadZipFile,KeyError): raise ValueError('El archivo no es un XLSX válido.')
    elif nombre.lower().endswith('.csv'):
        # Mismo formato ancho de dos encabezados; el nombre es el bimestre.
        from types import SimpleNamespace
        rows=list(csv.reader(io.StringIO(data.decode('utf-8-sig'))))
        if len(rows)>1000 or max(map(len,rows),default=0)>2000: raise ValueError('CSV demasiado grande.')
        width=max(map(len,rows),default=0)
        sheets=[(Path(nombre).stem,[[SimpleNamespace(value=v,number_format='') for v in row+['']*(width-len(row))] for row in rows])]
    else: raise ValueError('Solo se acepta CSV UTF-8 o XLSX sin macros.')
    facts, issues, dic, perfiles = [], [], {}, []
    def issue(code,detail,**kwargs): issues.append({'codigo':code,'detalle':detail,**kwargs})
    seen=set()
    for name,rows in sheets:
        p,label=periodo(name)
        if len(rows)<3: raise ValueError('Cada hoja necesita dos encabezados y al menos una empresa.')
        cols=columnas([c.value for c in rows[0]],[c.value for c in rows[1]],overrides)
        ids=[c['columna']-1 for c in cols if c['accion']=='identificador']
        if len(ids)!=1: raise ValueError(f'{name}: se requiere una columna Compañía inequívoca.')
        profile={'bimestre':p,'hoja':label,'columnas':len(cols),'empresas':0,'filas_omitidas':0,'sin_encabezado':0,'pii_descartadas':sum(c['accion']=='descartar_pii' for c in cols)}
        for c in cols:
            if c.get('clave'): dic[c['clave']]={k:v for k,v in c.items() if k not in {'columna','duplicada'}}
        for ri,row in enumerate(rows[2:],3):
            empresa=texto(row[ids[0]].value)
            if not empresa:
                profile['filas_omitidas']+=1
                substantive=any(c.value is not None and canon(c.value) not in VACIOS for c in row)
                issue('fila_sin_empresa' if substantive else 'fila_vacia_final','Fila omitida sin identificador.',bimestre=p,fila=ri)
                continue
            if not re.fullmatch(r'COMPA[ÑN][IÍ]A\s+\d{1,4}',empresa.upper()): raise ValueError('Use únicamente identificadores anónimos COMPAÑIA n.')
            empresa='COMPAÑIA '+re.search(r'\d+',empresa)[0]
            if (empresa,p) in seen: raise ValueError(f'Empresa y bimestre duplicados: {empresa}, {p}.')
            seen.add((empresa,p));profile['empresas']+=1
            for c in cols:
                cell=row[c['columna']-1]
                if c['accion']=='sin_encabezado':
                    if cell.value is not None and canon(cell.value) not in VACIOS:
                        profile['sin_encabezado']+=1
                        issue('columna_sin_encabezado','Valor no importado: falta pregunta/subopción; requiere corregir la fuente.',empresa=empresa,bimestre=p,fila=ri,columna=c['columna'])
                    continue
                if c['accion']!='incluir': continue
                val,txt,state=limpiar(cell.value,c['unidad'],cell.number_format)
                if c.get('duplicada'): state='encabezado_ambiguo';val=None;txt=None
                facts.append({'empresa':empresa,'bimestre':p,'metrica':c['metrica'],'dimension':c['dimension'],
                    'valor':val,'texto':txt,'estado':state,'clave':c['clave'],'unidad':'porcentaje' if c['unidad']=='proporcion' else c['unidad'],
                    'categoria':c['categoria'],'grupo_100':c['grupo_100'],'hoja':label,'fila':ri,'columna':c['columna'],
                    'excluido':False})
        perfiles.append(profile)
    # No confundir composiciones en fracciones sin formato con puntos porcentuales.
    # El analista puede declarar unidad=proporcion en el diccionario y reimportar.
    composiciones=defaultdict(list)
    for f in facts:
        if f['grupo_100'] and f['unidad']=='porcentaje':composiciones[(f['empresa'],f['bimestre'],f['grupo_100'])].append(f)
    for group in composiciones.values():
        nums=[f['valor'] for f in group if f['valor'] is not None]
        if nums and max(nums)<=1 and any(0<x<1 for x in nums) and not any(f['estado']=='porcentaje_excel_convertido' for f in group):
            for f in group:
                if f['valor'] is not None:
                    f['texto']=str(f['valor']);f['valor']=None;f['estado']='escala_porcentaje_ambigua'
    fingerprint=data if nombre.lower().endswith('.xlsx') else data+Path(nombre).stem.upper().encode()
    return {'huella':hashlib.sha256(fingerprint).hexdigest(),'hechos':facts,'incidencias':issues,'diccionario':dic,'perfil':perfiles}

def cargar_diccionario():
    p=RAIZ/'configuracion/diccionario_columnas.json'
    return json.loads(p.read_text()) if p.exists() else {}
