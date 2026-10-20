"""Recorridos completos y regresiones de los fallos encontrados en la auditoría."""
import csv
import io
import json
import os
import subprocess
import sys
from copy import deepcopy
from datetime import timedelta
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import select
from aplicacion.importador import columnas,leer_archivo,limpiar
from aplicacion.metricas import desglose,validar
from aplicacion.modelos import Hecho,Importacion
from conftest import csv_bytes,QUESTIONS,DIMS


def upload(client,data,name,dictionary=None):
    return client.post('/api/importar/vista-previa',files={'archivo':(name,data)},
                       data={'diccionario':json.dumps(dictionary)} if dictionary is not None else {})


def test_diccionario_corregido_reimporta_y_no_duplica(client,app):
    dictionary=client.get('/api/diccionario').json()
    key=next(k for k,v in dictionary.items() if v['metrica']=='plantilla_operativa')
    dictionary[key]['metrica']='plantilla_operativa_verificada'
    before=client.get('/api/filtros').json()['importacion']
    preview=upload(client,csv_bytes(),'JUL-AGO_2025.csv',dictionary)
    assert preview.status_code==200,preview.text
    assert client.get('/api/filtros').json()['importacion']==before
    activated=client.post('/api/importar/confirmar',json={'token':preview.json()['token']})
    assert activated.status_code==200,activated.text
    assert client.get('/api/filtros').json()['importacion']!=before
    assert client.get('/api/explorar?metrica=plantilla_operativa_verificada').json()['total']==7
    assert client.get('/api/panel').json()['serie'][0]['rotacion'] is None
    again=upload(client,csv_bytes(),'JUL-AGO_2025.csv',dictionary)
    assert client.post('/api/importar/confirmar',json={'token':again.json()['token']}).status_code==409
    with app.state.Session() as db:
        assert len(list(db.scalars(select(Importacion))))==2
        assert len(list(db.scalars(select(Importacion).where(Importacion.activa.is_(True)))))==1


def test_huella_antigua_rechaza_duplicado_sin_romper_compatibilidad(client,app,bundle):
    with app.state.Session.begin() as db:
        db.scalar(select(Importacion)).huella=bundle['huella_fuente']
    p=upload(client,csv_bytes(),'JUL-AGO_2025.csv')
    assert client.post('/api/importar/confirmar',json={'token':p.json()['token']}).status_code==409


@pytest.mark.parametrize('field,value',[('unidad','unidad_inventada'),('metrica',''),('accion','borrar'),('dimension',5)])
def test_diccionario_invalido_se_rechaza_explicitamente(client,field,value):
    key=next(c['clave'] for c in columnas(QUESTIONS,DIMS) if c.get('metrica')=='bajas')
    response=upload(client,csv_bytes(),'JUL-AGO_2025.csv',{key:{field:value}})
    assert response.status_code==422


def test_columnas_csv_vacias_no_heredan_pregunta():
    out=io.StringIO();csv.writer(out).writerows([
        ['Compañía','TOTAL DE CONTRATACIONES',''],['Open-Ended Response','Open-Ended Response',''],['COMPAÑIA 1',12,99]])
    result=leer_archivo(out.getvalue().encode(),'JUL-AGO_2025.csv')
    assert len(result['hechos'])==1 and result['hechos'][0]['valor']==12
    assert result['perfil'][0]['sin_encabezado']==1


def test_archivo_sin_empresas_o_metricas_no_activa(client):
    before=client.get('/api/filtros').json()['importacion']
    for data in [b'Compania,TOTAL DE CONTRATACIONES\nOpen-Ended Response,Open-Ended Response\n,\n',
                 b'Compania\nOpen-Ended Response\nCOMPANIA 1\n']:
        assert upload(client,data,'JUL-AGO_2025.csv').status_code==422
    assert client.get('/api/filtros').json()['importacion']==before


def test_desglose_no_mezcla_unidades():
    facts=[{'empresa':'COMPAÑIA 1','metrica':'duracion','dimension':'total','unidad':u,'valor':n,
            'estado':'valido','texto':None,'excluido':False} for u,n in [('horas',24),('dias',1)]]
    rows=desglose(facts,'duracion')
    assert {(r['unidad'],r['suma']) for r in rows}=={('horas',24),('dias',1)}


def test_horas_excel_tipadas_respetan_rangos():
    assert limpiar(timedelta(hours=-1),'horas')[2]=='negativo'
    assert limpiar(timedelta(hours=25),'hora_del_dia')[2]=='hora_fuera_rango'
    assert limpiar(timedelta(hours=25),'horas')==(25,None,'hora_convertida')


def test_atipico_temporal_visible_al_filtrar_una_empresa(bundle):
    facts=[]
    for p,n in [('2025-01',100),('2025-03',10000),('2025-05',100)]:
        fs=[deepcopy(f) for f in bundle['hechos'] if f['empresa']=='COMPAÑIA 1']
        for f in fs:
            f['bimestre']=p
            if f['metrica']=='plantilla_operativa':f['valor']=n
        facts.extend(fs)
    out=[x for x in validar(facts) if x['codigo']=='atipico_plantilla' and x['bimestre']=='2025-03']
    assert len(out)==1 and len(out[0]['historial'])==3


def test_xlsx_dos_esquemas_hasta_panel_exportacion_y_calidad(client,app):
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    for period,reverse in [('JUL-AGO_2024',False),('SEP-OCT_2024',True)]:
        sheet=workbook.create_sheet(period)
        rows=list(csv.reader(io.StringIO(csv_bytes().decode())))
        rows[0].append('Indicador nuevo del cuestionario');rows[1].append('Respuesta')
        for row in rows[2:]:row.append('Dato de prueba')
        if reverse:rows=[list(reversed(row)) for row in rows]
        for row in rows:sheet.append(row)
        sheet.append([None]*len(rows[0]));sheet.cell(len(rows)+1,1).number_format='0.00'
    content=io.BytesIO();workbook.save(content);workbook.close()
    response=upload(client,content.getvalue(),'estudio.xlsx');assert response.status_code==200
    preview=response.json();assert len(preview['perfil'])==2
    assert client.post('/api/importar/confirmar',json={'token':preview['token']}).status_code==200
    panel=client.get('/api/panel').json();assert len(panel['detalle'])==14
    assert [p['rotacion'] for p in panel['serie']]==[10,10]
    assert panel['serie'][1]['anterior']['n_pareadas']==7
    catalog=client.get('/api/catalogo').json()['metricas']
    assert sum(x['celdas'] for x in catalog)==preview['hechos']
    assert any(x['titulo']=='Indicador nuevo del cuestionario' for x in catalog)
    quality=client.get('/api/calidad').json();assert quality['conteos']['fila_vacia_final']==2
    export=client.get('/api/exportar');assert export.status_code==200
    assert len(list(csv.reader(io.StringIO(export.text))))==preview['hechos']+1
    assert 'PERSONA PRIVADA' not in export.text and 'secreto@example.com' not in export.text
    assert client.get('/api/reporte').status_code==200
    with app.state.Session() as db:
        facts=list(db.scalars(select(Hecho).where(Hecho.importacion_id==client.get('/api/filtros').json()['importacion'])))
        assert len(facts)==preview['hechos']


def test_cli_perfilar_e_importar_en_directorio_aislado(tmp_path):
    root=Path(__file__).resolve().parents[1];source=tmp_path/'JUL-AGO_2025.csv';source.write_bytes(csv_bytes())
    env={**os.environ,'PYTHONPATH':str(root),'RH_DATABASE_URL':'sqlite:///'+str(tmp_path/'estudio.db')}
    for command in ['perfilar','importar']:
        result=subprocess.run([sys.executable,'-m','aplicacion.cli',command,str(source)],cwd=tmp_path,env=env,capture_output=True,text=True)
        assert result.returncode==0,result.stderr
    quality=json.loads((tmp_path/'privado/perfil/calidad.json').read_text())
    assert quality['serie'][0]['rotacion']==10
    assert (tmp_path/'estudio.db').exists()
    assert 'secreto@example.com' not in (tmp_path/'privado/perfil/modelo_largo.jsonl').read_text()
